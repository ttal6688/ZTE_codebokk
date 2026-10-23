"""Build matched Markdown and DOCX explanations from audited minimal-run CSVs."""
from pathlib import Path
import json
import math
import re
from PIL import Image, ImageDraw, ImageFont
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'results/minimal_all_methods_20261022'
TITLE = 'UPA双极化双流预编码方法扩展说明'
STEM = '双流双极化全方法最小验证与扩展说明'
PAIRS = pd.read_csv(OUT / 'paired_summary.csv')
AUDIT = json.loads((OUT / 'complete_audit.json').read_text(encoding='utf-8'))
assert AUDIT['status'] == 'passed'
BLOCKS = []


def para(text): BLOCKS.append(('p', text))
def chapter(title): BLOCKS.append(('chapter', title))
def table(title, headers, rows, widths=None): BLOCKS.append(('table', (title, headers, rows, widths)))
def equation(latex, parts): BLOCKS.append(('eq', (latex, parts)))


def pipeline():
    assets=OUT/'assets';assets.mkdir(exist_ok=True)
    im=Image.new('RGB',(1600,400),'white');d=ImageDraw.Draw(im)
    font=ImageFont.truetype('C:/Windows/Fonts/simhei.ttf',34)
    small=ImageFont.truetype('C:/Windows/Fonts/simsun.ttc',28)
    specs=[(20,430,'单极化单流',['一个独立符号 s','一组极化端口','一个复系数']),
           (595,1005,'双极化单流',['同一个独立符号 s','两组极化端口','2×1 系数向量']),
           (1180,1590,'双极化双流',['两个独立符号 s1 s2','每组端口都可承载两流','2×2 系数矩阵'])]
    for left,right,title,lines in specs:
        d.rounded_rectangle((left,65,right,360),14,fill='#F1F5F8',outline='#36566F',width=3)
        for text,y,ft in [(title,93,font)]+[(line,165+i*51,small) for i,line in enumerate(lines)]:
            bb=d.textbbox((0,0),text,font=ft)
            d.text(((left+right-bb[2]+bb[0])/2,y),text,font=ft,fill='black')
    for x1,x2,label in [(440,585,'加极化端口'),(1015,1170,'加独立流')]:
        y=215;d.line((x1,y,x2,y),fill='#36566F',width=4)
        d.polygon([(x2,y),(x2-16,y-9),(x2-16,y+9)],fill='#36566F')
        bb=d.textbbox((0,0),label,font=small)
        d.text(((x1+x2-bb[2])/2,20),label,font=small,fill='black')
    im.save(assets/'三阶段扩展.png')


def build_content():
    global EXAMPLE_AUDIT
    from method_explanation_content import build
    EXAMPLE_AUDIT=build(para,chapter,table,equation,BLOCKS,PAIRS,AUDIT)


def report_scope():
    trials=pd.read_csv(OUT/'confirm/trials.csv')
    trials=trials[trials.array.eq('UPA')]
    assert len(trials)==2048
    result=dict(status='passed',array='UPA',records=len(trials),samples=32,
                channels=['LOS','MIX'],APs=8,observation_SNR_dB=30,
                reused_existing_data=True,branches={})
    for name,branches,expected in [('UE',['UE','UE_shared'],1024),('CPU',['CPU'],768)]:
        t=trials[trials.branch.isin(branches)];dual=t[t.streams.eq(2)]
        paired=PAIRS[PAIRS.array.eq('UPA')&PAIRS.branch.isin(branches)]
        assert len(t)==expected and len(dual)==expected//2
        assert paired.positive_count.eq(32).all() and dual.received_mode_ratio.gt(1e-12).all()
        assert (t.precoder_norm_squared-1).abs().max()<1e-12
        assert t.per_AP_power_error.max()<1e-12
        result['branches'][name]=dict(records=len(t),dual_records=len(dual),
            all_cases_positive_at_32_positions=True,all_dual_received_rank_two=True,
            minimum_received_mode_ratio=float(dual.received_mode_ratio.min()),
            maximum_total_power_error=float((t.precoder_norm_squared-1).abs().max()),
            maximum_per_AP_power_error=float(t.per_AP_power_error.max()),
            dual_feedback_bits={m:int(z.dual_bits.iloc[0]) for m,z in paired.groupby('method')})
    result['ideal_reference_records']=int(trials.branch.eq('ideal').sum())
    (OUT/'report_scope_UPA.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')


def md_table(headers, rows):
    return '\n'.join(['|'+'|'.join(headers)+'|','|'+'|'.join(['---']*len(headers))+'|']+['|'+'|'.join(map(str,row))+'|' for row in rows])


def markdown_equation(latex):
    # Keep matrix row separators apart from the next token. Some Markdown
    # viewers otherwise mistake the second slash plus j/c/s for a macro.
    formatted=latex.replace('\\\\', '\\\\\n')
    allowed={'Sigma','alpha','begin','beta','boldsymbol','end','frac','in',
             'ldots','mathbb','mathbf','mathrm','otimes','quad','sqrt','sum',
             'times','widehat','widetilde'}
    unknown=set(re.findall(r'\\([A-Za-z]+)',formatted))-allowed
    assert not unknown, f'Unsupported Markdown math macros: {sorted(unknown)}'
    return formatted


def markdown_content():
    md=['# '+TITLE,'','扩展原理与分链路最小验证  2026年10月22日','',f'[下载可转发的Word报告]({STEM}.docx)','']
    for kind,data in BLOCKS:
        if kind=='chapter':md.extend(['## '+data,''])
        elif kind=='p':md.extend([data,''])
        elif kind=='table':
            caption,heads,rows,_=data;md.extend(['**'+caption+'**','',md_table(heads,rows),''])
        elif kind=='eq':md.extend(['$$',markdown_equation(data[0]),'$$',''])
        elif kind=='image':
            path,caption=data;md.extend([f'![{caption}]({path})',''])
    return '\n'.join(md)


def set_font(run, name='宋体', size=11, bold=False):
    run.font.name=name;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    rpr=run._element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
    for key in ['eastAsia','ascii','hAnsi']:fonts.set(qn('w:'+key),name)


def native_equation(doc, parts):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(6)
    obj=OxmlElement('m:oMath')
    def fill(parent,value):
        for item in value if isinstance(value,list) else [value]:
            parent.append(node(item))
        return parent
    def node(part):
        if not isinstance(part,tuple):
            if len(str(part))==2 and str(part)[1] in ['̃','̂']:
                n=OxmlElement('m:acc');pr=OxmlElement('m:accPr');ch=OxmlElement('m:chr')
                ch.set(qn('m:val'),'̃' if str(part)[1]=='̃' else '̂');pr.append(ch);n.append(pr)
                n.append(fill(OxmlElement('m:e'),str(part)[0]));return n
            r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=str(part);r.append(t);return r
        kind=part[0]
        if kind in ['sub','sup','sum']:
            n=OxmlElement('m:'+{'sub':'sSub','sup':'sSup','sum':'sSubSup'}[kind])
            n.append(fill(OxmlElement('m:e'),'∑' if kind=='sum' else part[1]))
            if kind in ['sub','sum']:n.append(fill(OxmlElement('m:sub'),part[1] if kind=='sum' else part[2]))
            if kind in ['sup','sum']:n.append(fill(OxmlElement('m:sup'),part[2]))
            return n
        if kind=='frac':
            n=OxmlElement('m:f');n.append(fill(OxmlElement('m:num'),part[1]));n.append(fill(OxmlElement('m:den'),part[2]));return n
        if kind=='mat':
            delim=OxmlElement('m:d');pr=OxmlElement('m:dPr')
            for name,val in [('begChr','['),('endChr',']')]:
                el=OxmlElement('m:'+name);el.set(qn('m:val'),val);pr.append(el)
            delim.append(pr);base=OxmlElement('m:e');matrix=OxmlElement('m:m')
            for row in part[1]:
                mr=OxmlElement('m:mr')
                for cell in row:mr.append(fill(OxmlElement('m:e'),cell))
                matrix.append(mr)
            base.append(matrix);delim.append(base);return delim
        raise ValueError(kind)
    fill(obj,parts);p._p.append(obj)


def word_table(doc, title, headers, rows, widths):
    cap=doc.add_paragraph();cap.paragraph_format.keep_with_next=True;cap.paragraph_format.space_before=Pt(6);cap.paragraph_format.space_after=Pt(5);set_font(cap.add_run(title),size=10.5,bold=True)
    tb=doc.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    if widths is None:widths=[16.4/len(headers)]*len(headers)
    for col,width in zip(tb.columns,widths):col.width=Cm(width)
    for idx,values in enumerate([headers]+rows):
        cells=tb.rows[0].cells if idx==0 else tb.add_row().cells
        for j,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders);margin=OxmlElement('w:tcMar')
            for edge,valuepad in [('top','80'),('bottom','80'),('left','90'),('right','90')]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),valuepad);e.set(qn('w:type'),'dxa');margin.append(e)
            tcpr.append(margin)
            if idx==0 or idx%2==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF4' if idx==0 else 'F7F9FB');tcpr.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.10
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if idx==0 or (j>0 and len(headers)>=4) else WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(str(value)),size=10 if len(headers)>=4 else 10.5,bold=idx==0)
        pr=tb.rows[idx]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if idx==0:pr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)


def write_reports():
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    # Remove built-in title/subtitle rules inherited from the default template.
    for node in doc.styles.element.xpath('.//w:pBdr'):
        node.getparent().remove(node)
    sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
    sec.footer_distance=Cm(.8)
    for name,size in [('Normal',11),('Title',20),('Heading 1',15)]:
        st=doc.styles[name];st.font.name='宋体' if name=='Normal' else '黑体';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().append(OxmlElement('w:rFonts'));st.element.rPr.rFonts.set(qn('w:eastAsia'),st.font.name)
    normal=doc.styles['Normal'].paragraph_format;normal.line_spacing=1.16;normal.space_after=Pt(5)
    title=doc.add_paragraph(style='Title');set_font(title.add_run(TITLE),'黑体',20)
    title.paragraph_format.space_after=Pt(12)
    p=doc.add_paragraph();set_font(p.add_run('扩展原理与分链路最小验证    2026年10月22日'),size=10.5)
    p.paragraph_format.space_after=Pt(14)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    for kind,data in BLOCKS:
        if kind=='chapter':
            doc.add_page_break();p=doc.add_paragraph(data,style='Heading 1');p.paragraph_format.space_after=Pt(10)
        elif kind=='p':
            doc.add_paragraph(data)
        elif kind=='table':
            caption,heads,rows,widths=data;word_table(doc,caption,heads,rows,widths)
        elif kind=='eq':
            latex,parts=data;native_equation(doc,parts)
        elif kind=='image':
            path,caption=data;p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run();r.add_picture(str(OUT/path),width=Cm(16.0))
            pr=r._r.xpath('.//wp:docPr')[0];pr.set('descr',caption)
            p=doc.add_paragraph(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.core_properties.title=TITLE;doc.core_properties.subject='UPA双极化双流扩展原理及UE和CPU独立验证';doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.save(OUT/(STEM+'.docx'));(OUT/(STEM+'.md')).write_text(markdown_content(),encoding='utf-8')
    (OUT/'teaching_example_audit.json').write_text(json.dumps(EXAMPLE_AUDIT,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Reports written:',STEM)


if __name__=='__main__':
    import sys
    if '--markdown-only' in sys.argv:
        build_content();(OUT/(STEM+'.md')).write_text(markdown_content(),encoding='utf-8')
        print('Markdown updated; existing Word file preserved')
    else:
        pipeline();build_content();report_scope();write_reports()
