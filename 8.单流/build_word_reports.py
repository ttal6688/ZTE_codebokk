from pathlib import Path
import re,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent.parent/'8.单流'
DOCS=ROOT/'docs'

def mr(text):
    node=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;node.append(t);return node

def sub(text,index):
    node=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(text));a=OxmlElement('m:sub');a.append(mr(index));node.extend([e,a]);return node

def equation(doc,formula):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    container=OxmlElement('m:oMathPara');o=OxmlElement('m:oMath');container.append(o)
    if formula.startswith('W='):o.extend([mr('W = F C')])
    elif formula.startswith('C_'):
        two='2B' in formula;index='b,s' if two else 'b,1';o.append(sub('C',index));o.append(mr(' = '))
        frac=OxmlElement('m:f');num=OxmlElement('m:num');den=OxmlElement('m:den')
        power=OxmlElement('m:sSup');base=OxmlElement('m:e');base.append(mr('e'));ex=OxmlElement('m:sup');ex.extend([mr('j'),sub('θ',index)]);power.extend([base,ex]);num.append(power)
        root=OxmlElement('m:rad');prop=OxmlElement('m:radPr');hide=OxmlElement('m:degHide');hide.set(qn('m:val'),'1');prop.append(hide);root.append(prop);root.append(OxmlElement('m:deg'));value=OxmlElement('m:e');value.append(mr('2B' if two else 'B'));root.append(value);den.append(root);frac.extend([num,den]);o.append(frac)
        if two:o.append(mr(' ,  s = 1, 2'))
    else:
        o.append(mr('R = '));summ=OxmlElement('m:nary');prop=OxmlElement('m:naryPr');char=OxmlElement('m:chr');char.set(qn('m:val'),'∑');prop.append(char);lim=OxmlElement('m:limLoc');lim.set(qn('m:val'),'undOvr');prop.append(lim);summ.append(prop)
        a=OxmlElement('m:sub');a.append(mr('s=1'));b=OxmlElement('m:sup');b.append(mr('2'));e=OxmlElement('m:e');e.extend([sub('log','2'),mr('(1 + '),sub('γ','s'),mr(')')]);summ.extend([a,b,e]);o.append(summ)
    p._p.append(container);p.paragraph_format.space_after=Pt(3)

def plain(text):
    text=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',text)
    return text.replace('**','').replace('`','')

def build(md):
    doc=Document();section=doc.sections[0];section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.78);section.left_margin=section.right_margin=Inches(.85)
    for name in ['Normal','Body Text','List Bullet','List Number']:
        style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(11)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
        style.paragraph_format.space_after=Pt(6);style.paragraph_format.line_spacing=1.12
    for name,size in [('Title',22),('Heading 1',15),('Heading 2',12.5)]:
        style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑');style.paragraph_format.space_before=Pt(12);style.paragraph_format.space_after=Pt(7)
        style.paragraph_format.keep_with_next=True
    # Remove inherited template title/heading rules; keep all titles black.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run('第 ');run.font.size=Pt(9);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')
    lines=md.read_text(encoding='utf-8').splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line=='$$':
            formula=[];i+=1
            while i<len(lines) and lines[i].strip()!='$$':formula.append(lines[i]);i+=1
            equation(doc,''.join(formula));i+=1;continue
        if line.startswith('```'):
            code=[];i+=1
            while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
            for text in code:
                p=doc.add_paragraph(text);p.paragraph_format.space_after=Pt(3)
                for run in p.runs:run.font.name='Consolas';run.font.size=Pt(9)
            i+=1;continue
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[plain(x.strip()) for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',x) for x in row):block.append(row)
                i+=1
            n=len(block[0]);
            if doc.paragraphs and doc.paragraphs[-1].text:doc.paragraphs[-1].paragraph_format.keep_with_next=True
            table=doc.add_table(rows=0,cols=n);table.style='Table Grid';table.autofit=False
            widths=([1.6,5.2] if n==2 else ([1.65]+[(6.8-1.65)/(n-1)]*(n-1)))
            for rowidx,values in enumerate(block):
                cells=table.add_row().cells
                for j,(cell,text) in enumerate(zip(cells,values)):
                    cell.width=Inches(widths[j]);cell.text=text
                    tcPr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
                    for edge,v in [('top','50'),('bottom','50'),('left','80'),('right','80')]:
                        x=OxmlElement('w:'+edge);x.set(qn('w:w'),v);x.set(qn('w:type'),'dxa');margins.append(x)
                    tcPr.append(margins)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
                        p.paragraph_format.keep_with_next=(rowidx<len(block)-1)
                        if j>0 and n>2:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                        for run in p.runs:run.font.size=Pt(10.5);run.bold=(rowidx==0)
                cant=OxmlElement('w:cantSplit');table.rows[-1]._tr.get_or_add_trPr().append(cant)
                if rowidx==0:
                    header=OxmlElement('w:tblHeader');table.rows[-1]._tr.get_or_add_trPr().append(header)
            doc.add_paragraph().paragraph_format.space_after=Pt(1)
            continue
        match=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)',line)
        if match:
            image=(md.parent/match.group(2)).resolve();assert image.exists()
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(image),width=Inches(6.0))
            caption=doc.add_paragraph(match.group(1));caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_after=Pt(9)
            for run in caption.runs:run.font.size=Pt(10)
        elif line.startswith('# '):doc.add_paragraph(plain(line[2:]),'Title')
        elif line.startswith('## '):doc.add_paragraph(plain(line[3:]),'Heading 1')
        elif line.startswith('### '):doc.add_paragraph(plain(line[4:]),'Heading 2')
        elif line.startswith('- '):doc.add_paragraph(plain(line[2:]),'List Bullet')
        else:doc.add_paragraph(plain(line))
        i+=1
    dest=md.with_suffix('.docx');doc.save(dest);return dest

if __name__=='__main__':
    for name in ['UPA单流400位置正式实验报告.md','UPA单极化双流扩展验证报告.md']:
        print(build(DOCS/name))
